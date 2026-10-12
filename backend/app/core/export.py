import io
import pandas as pd
from typing import List, Dict
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


def create_excel_report(ranked_results: List[Dict], job_description: str) -> bytes:
    """
    Generate a professional Excel report from ranked resume results.
    Returns the Excel file as bytes.
    """
    rows = []
    for item in ranked_results:
        rows.append({
            "Rank": item.get("rank"),
            "Filename": item.get("filename"),
            "Score": item.get("score"),
            "Matched Skills": ", ".join(item.get("skill_match", {}).get("matched_skills", item.get("matched_skills", []))),
            "Missing Skills": ", ".join(item.get("skill_match", {}).get("missing_skills", item.get("missing_skills", []))),
            "Skill Match %": item.get("skill_match", {}).get("match_percentage", ""),
        })

    df = pd.DataFrame(rows)

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Ranked Resumes")

        # Optional second sheet with JD preview
        jd_df = pd.DataFrame({"Job Description Preview": [job_description[:2000]]})
        jd_df.to_excel(writer, index=False, sheet_name="Job Description")

        workbook = writer.book
        worksheet = writer.sheets["Ranked Resumes"]

        # Styling
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        for col_num, column_title in enumerate(df.columns, 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border

        # Auto column width
        for col_num, column_title in enumerate(df.columns, 1):
            max_length = max(
                len(str(column_title)),
                *(len(str(df.iloc[i, col_num - 1])) for i in range(min(len(df), 50)))
            )
            worksheet.column_dimensions[get_column_letter(col_num)].width = min(max_length + 4, 50)

        # Color top 3 ranks
        green_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        yellow_fill = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")
        orange_fill = PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid")

        for row_idx in range(2, len(df) + 2):
            rank_cell = worksheet.cell(row=row_idx, column=1)
            rank = rank_cell.value
            if rank == 1:
                for col in range(1, len(df.columns) + 1):
                    worksheet.cell(row=row_idx, column=col).fill = green_fill
            elif rank == 2:
                for col in range(1, len(df.columns) + 1):
                    worksheet.cell(row=row_idx, column=col).fill = yellow_fill
            elif rank == 3:
                for col in range(1, len(df.columns) + 1):
                    worksheet.cell(row=row_idx, column=col).fill = orange_fill

    output.seek(0)
    return output.read()